#!/usr/bin/env python3
"""Tidy table pagination in a pandoc-built .docx.

Two fixes, applied to every table:
  1. cantSplit on every row  — a row never splits across a page boundary.
  2. keepNext on header-row cell paragraphs — the (repeating) header row binds to
     the first body row, so it can never be stranded alone at the bottom of a page
     leaving a near-empty page (the artefact we saw on the agents table).

Run AFTER pandoc and BEFORE add_cover.py (so the cover injection is untouched):
  python3 scripts/fix-table-breaks.py <book.docx>
"""
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def ensure_child(parent, tag):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    return el


def main():
    path = sys.argv[1]
    doc = Document(path)
    rows_fixed = 0
    headers_fixed = 0
    for table in doc.tables:
        for i, row in enumerate(table.rows):
            tr = row._tr
            trPr = tr.get_or_add_trPr()
            # 1) never split a row across pages
            ensure_child(trPr, 'w:cantSplit')
            rows_fixed += 1
            # 2) header row (row 0, or any row pandoc marked to repeat): keep it
            #    bound to the following row so it can't be stranded on its own page.
            is_header = i == 0 or trPr.find(qn('w:tblHeader')) is not None
            if is_header:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        pPr = p._p.get_or_add_pPr()
                        ensure_child(pPr, 'w:keepNext')
                headers_fixed += 1
    doc.save(path)
    print(f"  table-break fix: {rows_fixed} rows cantSplit, {headers_fixed} header rows keepNext")


if __name__ == '__main__':
    main()
